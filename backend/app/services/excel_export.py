"""把一场考试的成绩导出成 Excel，供教师登分、存档、发教研组。"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="F0F2F5")
LOW_SCORE_FILL = PatternFill("solid", fgColor="FFF1F0")  # 得分率低于60%的格子标红


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _autosize(ws, min_width: int = 8, max_width: int = 40) -> None:
    for column in ws.columns:
        longest = max((len(str(c.value or "")) for c in column), default=0)
        letter = get_column_letter(column[0].column)
        # 中文字符实际显示宽度约为英文的两倍，粗略乘 1.6 够用
        ws.column_dimensions[letter].width = max(min_width, min(max_width, int(longest * 1.6) + 2))


def build_workbook(stats: dict, submissions: list) -> bytes:
    wb = Workbook()

    # ---- 成绩表：一人一行，后面跟着逐题得分 ----
    ws = wb.active
    ws.title = "成绩表"

    questions = stats.get("questions") or []
    headers = ["名次", "学生", "总分", "得分率"] + [
        f"{q['question_no']}({q['max_score']})" for q in questions
    ]
    if stats.get("essay"):
        headers.append(f"作文({stats['essay']['max_score']})")
    ws.append(headers)

    by_id = {s.id: s for s in submissions}
    for rank, student in enumerate(stats.get("students") or [], 1):
        sub = by_id.get(student["submission_id"])
        result = (sub.result or {}) if sub else {}
        row = [
            rank,
            student["name"],
            student["score"],
            round((student.get("rate") or 0) * 100, 1) / 100,
        ]
        # 逐题按位置对齐——题号写法在不同试卷间并不稳定
        items = result.get("questions") or []
        for i in range(len(questions)):
            row.append(items[i].get("score") if i < len(items) else None)
        if stats.get("essay"):
            row.append((result.get("essay") or {}).get("score"))
        ws.append(row)

    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            cell.number_format = "0.0%"
    _style_header(ws)
    _autosize(ws)

    # ---- 逐题分析：哪道题最该讲评 ----
    ws2 = wb.create_sheet("逐题分析")
    ws2.append(["题号", "大题", "满分", "平均分", "得分率", "满分人数", "零分人数", "主要失分原因"])
    for q in sorted(questions, key=lambda x: x.get("score_rate") if x.get("score_rate") is not None else 1):
        ws2.append(
            [
                q["question_no"],
                q.get("section") or "",
                q["max_score"],
                q["avg_score"],
                q.get("score_rate"),
                q["full_marks"],
                q["zero_marks"],
                "；".join(q.get("common_reasons") or []),
            ]
        )
    for row in ws2.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "0.0%"
            if isinstance(cell.value, (int, float)) and cell.value < 0.6:
                cell.fill = LOW_SCORE_FILL
    _style_header(ws2)
    _autosize(ws2)
    ws2.column_dimensions["H"].width = 60

    # ---- 总体情况 ----
    ws3 = wb.create_sheet("总体情况")
    rows = [
        ("考试名称", stats.get("exam_name")),
        ("试卷满分", stats.get("total_score")),
        ("已批改", f"{stats.get('graded_count')} / {stats.get('submission_count')}"),
        ("平均分", stats.get("avg_score")),
        ("中位数", stats.get("median_score")),
        ("最高分", stats.get("max_score_got")),
        ("最低分", stats.get("min_score_got")),
        ("及格率", stats.get("pass_rate")),
        ("优秀率", stats.get("excellent_rate")),
    ]
    ws3.append(["项目", "数值"])
    for name, value in rows:
        ws3.append([name, value])
    for name in ("及格率", "优秀率"):
        for row in ws3.iter_rows(min_row=2):
            if row[0].value == name:
                row[1].number_format = "0.0%"
    ws3.append([])
    ws3.append(["分数段", "人数"])
    for band, count in (stats.get("bands") or {}).items():
        ws3.append([band, count])
    _style_header(ws3)
    _autosize(ws3)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
