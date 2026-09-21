"""参考答案文本提取（工作流文档第5章：不能直接拿整篇原文模糊评分，先取出原文供人工/AI结构化）。"""

from pathlib import Path

from docx import Document


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")

    raise ValueError(f"暂不支持解析 {suffix} 格式的参考答案，请先转换为 Word(.docx) 或文本文件")
